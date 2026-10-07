import json
import os
import tkinter as tk
from tkinter import filedialog, messagebox

import openpyxl


def split_values(value):
    """Transforme 'a; b; c' en ['a', 'b', 'c']."""
    if value is None or str(value).strip() == "":
        return []

    return [
        item.strip()
        for item in str(value).split(";")
        if item.strip()
    ]


def clean_value(value):
    """Nettoie une cellule Excel."""
    if value is None:
        return None

    if isinstance(value, str):
        value = value.strip()

        if value == "":
            return None

    return value


def export_cures(sheet):
    characters = []

    # Première ligne = noms des colonnes
    headers = [
        cell.value
        for cell in sheet[1]
    ]

    # Parcours des personnages
    for row in sheet.iter_rows(min_row=2, values_only=True):

        # Ignore les lignes complètement vides
        if not any(value is not None for value in row):
            continue

        data = dict(zip(headers, row))

        character = {
            "id": clean_value(data.get("Value.id")),
            "name": clean_value(data.get("Value.name")),
            "cure_name": clean_value(data.get("Value.cure_name")),

            "categories": {
                "season": split_values(
                    data.get("Value.categories.season")
                ),

                "generation": clean_value(
                    data.get("Value.categories.generation")
                ),

                "cure_hair_color": clean_value(
                    data.get("Value.categories.cure_hair_color")
                ),

                "cure_main_color": clean_value(
                    data.get("Value.categories.cure_main_color")
                ),

                "cure_eyes_color": clean_value(
                    data.get("Value.categories.cure_eyes_color")
                )
            },

            "image": clean_value(
                data.get("Value.image")
            )
        }

        characters.append(character)

    return {
        "characters": characters
    }


def export_films(sheet):
    films = []

    # Première ligne = noms des colonnes
    headers = [
        cell.value
        for cell in sheet[1]
    ]

    # Parcours des films
    for row in sheet.iter_rows(min_row=2, values_only=True):

        # Ignore les lignes complètement vides
        if not any(value is not None for value in row):
            continue

        data = dict(zip(headers, row))

        film = {
            "name": clean_value(
                data.get("Value.name")
            ),

            "characters": [
                int(value)
                for value in split_values(
                    data.get("Value.characters")
                )
            ],

            "type": clean_value(
                data.get("Value.type")
            )
        }

        films.append(film)

    return {
        "films": films
    }


def save_json(data, filename, folder):
    path = os.path.join(folder, filename)

    with open(path, "w", encoding="utf-8") as file:
        json.dump(
            data,
            file,
            ensure_ascii=False,
            indent=4
        )

    return path


def choose_excel():
    path = filedialog.askopenfilename(
        title="Choisir le fichier Excel",
        filetypes=[
            ("Fichiers Excel", "*.xlsx"),
            ("Tous les fichiers", "*.*")
        ]
    )

    if not path:
        return

    excel_path.set(path)


def generate_json():
    path = excel_path.get()

    if not path:
        messagebox.showwarning(
            "Fichier manquant",
            "Sélectionne d'abord ton fichier Excel."
        )
        return

    if not os.path.exists(path):
        messagebox.showerror(
            "Erreur",
            "Le fichier Excel sélectionné n'existe pas."
        )
        return

    try:
        workbook = openpyxl.load_workbook(
            path,
            data_only=True
        )

        # Vérification des feuilles
        if "Cures" not in workbook.sheetnames:
            messagebox.showerror(
                "Erreur",
                "La feuille 'Cures' est introuvable."
            )
            return

        if "Films" not in workbook.sheetnames:
            messagebox.showerror(
                "Erreur",
                "La feuille 'Films' est introuvable."
            )
            return

        cures_sheet = workbook["Cures"]
        films_sheet = workbook["Films"]

        cures_data = export_cures(cures_sheet)
        films_data = export_films(films_sheet)

        # Les JSON sont créés dans le même dossier que le XLSX
        folder = os.path.dirname(path)

        cures_path = save_json(
            cures_data,
            "Cures.json",
            folder
        )

        films_path = save_json(
            films_data,
            "Films.json",
            folder
        )

        messagebox.showinfo(
            "Export terminé",
            "Les fichiers ont été générés avec succès !\n\n"
            "Cures.json\n"
            "Films.json\n\n"
            f"Dossier :\n{folder}"
        )

    except Exception as error:
        messagebox.showerror(
            "Erreur",
            f"Une erreur est survenue :\n\n{error}"
        )


# =========================================================
# INTERFACE
# =========================================================

root = tk.Tk()
root.title("DLE - Exporteur JSON")
root.geometry("600x220")
root.resizable(False, False)

excel_path = tk.StringVar()

title = tk.Label(
    root,
    text="DLE - Exporteur JSON",
    font=("Arial", 18, "bold")
)

title.pack(pady=(20, 15))

description = tk.Label(
    root,
    text="Sélectionne ton fichier Database.xlsx puis génère les JSON.",
    font=("Arial", 10)
)

description.pack()

frame = tk.Frame(root)
frame.pack(pady=20)

entry = tk.Entry(
    frame,
    textvariable=excel_path,
    width=55
)

entry.grid(row=0, column=0, padx=(0, 8))

browse_button = tk.Button(
    frame,
    text="Parcourir...",
    command=choose_excel
)

browse_button.grid(row=0, column=1)

export_button = tk.Button(
    root,
    text="GÉNÉRER LES JSON",
    command=generate_json,
    font=("Arial", 11, "bold"),
    bg="#4CAF50",
    fg="white",
    padx=20,
    pady=8
)

export_button.pack()

root.mainloop()
